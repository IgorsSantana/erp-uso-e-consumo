"""Regressions for fresh deployments and the zero-limit catalog release."""
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import unittest

import openpyxl

from backend.app import create_app
from backend.bundled_limits import apply_limits, read_limits
from backend.db import SCHEMA
from backend.stock_limits import migrate_limits
from backend.workbook_limits import extract_limits


class BundledLimitTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix='erp-bundled-limits-')
        self.addCleanup(directory.cleanup)
        self.path = Path(directory.name)/'erp.sqlite3'
        self.config = {'TESTING': True, 'DATABASE': str(self.path), 'SECRET_KEY': 'seed-limit-test-only-key'}

    def connect(self):
        db = sqlite3.connect(self.path)
        db.row_factory = sqlite3.Row
        db.execute('PRAGMA foreign_keys=ON')
        self.addCleanup(db.close)
        return db

    def values(self, db, pid, sid):
        row = db.execute('SELECT estoque_minimo,estoque_maximo FROM estoque_limites WHERE produto_id=? AND loja_id=?', (pid, sid)).fetchone()
        return tuple(row) if row else None

    def test_fresh_start_populates_each_source_store_after_creating_catalog(self):
        app = create_app({**self.config, 'SEED_CATALOG': True})
        db = self.connect()
        source = read_limits()
        self.assertEqual(db.execute('SELECT COUNT(*) FROM produtos').fetchone()[0], 466)
        self.assertEqual(db.execute('SELECT COUNT(*) FROM estoque_limites').fetchone()[0], 629)
        for pair in source['limits']:
            sid = db.execute('SELECT id FROM lojas WHERE sigla=?', (pair['store'],)).fetchone()[0]
            self.assertEqual(self.values(db, pair['id'], sid), (pair['estoque_minimo'], pair['estoque_maximo']))
        self.assertEqual(self.values(db, '369842', 1), (20, 40))
        self.assertEqual(self.values(db, '369842', 3), (10, 20))
        self.assertEqual(self.values(db, '367054', 1), (0, 7))
        self.assertEqual(self.values(db, '367054', 3), (0, 2))
        self.assertEqual(self.values(db, '337086', 1), (1, 2))
        self.assertIsNone(self.values(db, '337086', 3))
        self.assertIsNone(self.values(db, '366959', 1))
        self.assertIsNone(self.values(db, '334512', 1))
        self.assertEqual(db.execute('SELECT COUNT(*) FROM estoque').fetchone()[0], 0)
        self.assertEqual(db.execute('SELECT COUNT(*) FROM movimentacoes').fetchone()[0], 0)
        client = app.test_client()
        login = client.post('/api/setup', json={'nome': 'Gestor', 'username': 'admin_seed', 'password': 'SenhaSeguroSeed2026!'})
        self.assertEqual(login.status_code, 201)
        limits = client.get('/api/produtos/369842/limites').json['lojas']
        self.assertEqual(next(item for item in limits if item['loja_id'] == 3)['estoque_maximo'], 20)
        alerts = client.get('/api/dashboard').json['alertas']
        self.assertTrue(any(item['produto_id'] == '369842' and item['loja_id'] == 3 for item in alerts))
        self.assertFalse(any(item['produto_id'] == '367054' for item in alerts))

    def test_previously_seeded_zero_database_is_repaired_without_reimport(self):
        create_app(self.config)
        db = self.connect()
        # Reproduce the previous startup: migration first, catalog second,
        # without the newly introduced per-store release.
        from backend.bundled_catalog import apply_catalog
        apply_catalog(db)
        db.execute("INSERT INTO estoque_limites(produto_id,loja_id) VALUES('369842',1)")
        db.commit()
        self.assertEqual(db.execute('SELECT COUNT(*) FROM catalog_releases').fetchone()[0], 1)
        self.assertEqual(self.values(db, '369842', 1), (0, 0))
        create_app({**self.config, 'SEED_CATALOG': True})
        self.assertEqual(self.values(db, '369842', 1), (20, 40))
        self.assertEqual(db.execute('SELECT COUNT(*) FROM catalog_releases').fetchone()[0], 1)
        self.assertEqual(db.execute('SELECT COUNT(*) FROM stock_limit_releases').fetchone()[0], 1)

    def legacy_database(self):
        db = self.connect()
        db.executescript(SCHEMA)
        db.execute('PRAGMA user_version=1')
        # Nonstandard IDs prove that store matching uses the registered sigla.
        db.executemany('INSERT INTO lojas(id,nome,sigla) VALUES(?,?,?)', [(41, 'Loja BCS', 'BCS'), (42, 'Loja SJN', 'SJN'), (43, 'Nova', 'NOVA')])
        db.execute("INSERT INTO produtos(id,nome,estoque_minimo,estoque_maximo) VALUES('369842','AGLOMAX',30,60)")
        db.execute("INSERT INTO produtos(id,nome,unidade,estoque_minimo,estoque_maximo) VALUES('337086','AÇAFRÃO','KG',1,2)")
        db.commit()
        migrate_limits(db)
        # Deterministic equivalent of the original migration timestamp.
        db.execute("UPDATE estoque_limites SET atualizado_em='2026-10-19T19:54:21Z'")
        db.execute("UPDATE schema_migrations SET applied_at='2026-10-19 19:54:21'")
        db.commit()
        return db

    def test_untouched_global_copies_are_replaced_only_for_known_source_stores(self):
        db = self.legacy_database()
        result = apply_limits(db)
        self.assertEqual(self.values(db, '369842', 41), (20, 40))
        self.assertEqual(self.values(db, '369842', 42), (10, 20))
        self.assertEqual(self.values(db, '337086', 41), (1, 2))
        self.assertIsNone(self.values(db, '337086', 42))
        self.assertEqual(self.values(db, '369842', 43), (30, 60))
        self.assertEqual(result['copias_globais_removidas'], 1)
        self.assertEqual(db.execute('PRAGMA foreign_key_check').fetchall(), [])

    def test_manual_positive_values_are_preserved_even_without_audit(self):
        db = self.legacy_database()
        db.execute("UPDATE estoque_limites SET estoque_minimo=3,estoque_maximo=9 WHERE produto_id='369842' AND loja_id=42")
        db.commit()
        apply_limits(db)
        self.assertEqual(self.values(db, '369842', 42), (3, 9))
        self.assertEqual(self.values(db, '369842', 41), (20, 40))

    def test_explicit_zero_and_removed_settings_are_preserved_using_audit(self):
        for cleared in (False, True):
            with self.subTest(cleared=cleared):
                # Recreate a clean DB between cases.
                if self.path.exists():
                    self.path.unlink()
                db = self.legacy_database()
                db.execute("DELETE FROM estoque_limites WHERE produto_id='369842' AND loja_id=42")
                if not cleared:
                    db.execute("INSERT INTO estoque_limites(produto_id,loja_id) VALUES('369842',42)")
                db.execute('INSERT INTO auditoria(acao,entidade,entidade_id,detalhes) VALUES(?,?,?,?)',
                           ('configurar_limites', 'produtos', '369842', json.dumps([{'loja_id': 42, 'configurado': not cleared}])))
                db.commit()
                apply_limits(db)
                self.assertEqual(self.values(db, '369842', 42), None if cleared else (0, 0))
                self.assertEqual(self.values(db, '369842', 41), (20, 40))
                db.close()

    def test_csv_limit_import_is_preserved_and_unit_mismatch_is_reported(self):
        db = self.legacy_database()
        db.execute("UPDATE produtos SET unidade='CX' WHERE id='369842'")
        db.execute('INSERT INTO auditoria(acao,entidade,entidade_id,detalhes) VALUES(?,?,?,?)',
                   ('importar', 'produtos', 'CSV', json.dumps({'loja_limites_id': 42, 'limites_atualizados': 1})))
        db.commit()
        result = apply_limits(db)
        self.assertEqual(self.values(db, '369842', 41), (30, 60))
        self.assertEqual(self.values(db, '369842', 42), (30, 60))
        self.assertIn('369842', result['unidades_para_revisao'])

    def test_restarts_preserve_edits_deletions_products_and_new_stores(self):
        create_app({**self.config, 'SEED_CATALOG': True})
        db = self.connect()
        db.execute("UPDATE estoque_limites SET estoque_minimo=5,estoque_maximo=15 WHERE produto_id='369842' AND loja_id=1")
        db.execute("DELETE FROM estoque_limites WHERE produto_id='369842' AND loja_id=3")
        db.execute("DELETE FROM produtos WHERE id='337086'")
        db.execute("INSERT INTO lojas(id,nome,sigla) VALUES(90,'Loja Nova','NOV')")
        db.commit()
        create_app({**self.config, 'SEED_CATALOG': True})
        self.assertEqual(self.values(db, '369842', 1), (5, 15))
        self.assertIsNone(self.values(db, '369842', 3))
        self.assertIsNone(self.values(db, '369842', 90))
        self.assertEqual(db.execute("SELECT COUNT(*) FROM produtos WHERE id='337086'").fetchone()[0], 0)
        self.assertFalse(apply_limits(db)['aplicado'])

    def test_failure_rolls_back_every_store_and_release_marker(self):
        db = self.legacy_database()
        before = [tuple(row) for row in db.execute('SELECT * FROM estoque_limites ORDER BY produto_id,loja_id')]
        db.execute("CREATE TRIGGER fail_limit_audit BEFORE INSERT ON auditoria WHEN NEW.acao='importar_limites_lojas' BEGIN SELECT RAISE(ABORT,'rollback'); END")
        db.commit()
        with self.assertRaises(sqlite3.IntegrityError):
            apply_limits(db)
        self.assertEqual([tuple(row) for row in db.execute('SELECT * FROM estoque_limites ORDER BY produto_id,loja_id')], before)
        self.assertEqual(db.execute('SELECT COUNT(*) FROM stock_limit_releases').fetchone()[0], 0)

    def test_invalid_release_never_changes_the_database(self):
        db = self.legacy_database()
        for value in (-1, 'NaN', 1.0001):
            with self.subTest(value=value):
                data = read_limits()
                data['limits'][0]['estoque_maximo'] = value
                path = self.path.parent/'invalid.json'
                path.write_text(json.dumps(data), encoding='utf-8')
                with self.assertRaises(ValueError):
                    apply_limits(db, path)
        self.assertEqual(self.values(db, '369842', 41), (30, 60))


class WorkbookStoreColumnTests(unittest.TestCase):
    def workbook(self):
        book = openpyxl.Workbook()
        book.remove(book.active)
        sheet = book.create_sheet('LIMPEZA')
        sheet.append(['BCS', 'SJN', 'CÓDIGO', 'DESCRIÇÃO PRODUTO', 'BCS ', 'SJN 2', 'MEP 2', 'FCL1 2', 'FCL2 2', 'FCL3 2', 'FCL4', 'ESTOQUE MÁXIMO EM GALÃO', 'ESTOQUE'])
        sheet.append(['X', 'X', 123, 'Limpador', 7, 2, 2, 1, 1, 1, 1, 15, 999])
        sheet = book.create_sheet('PADARIA')
        sheet.append(['CÓDIGO', 'DESCRIÇÃO PRODUTO', 'ESTOQUE MÍNIMO', 'ESTOQUE MÁXIMO', 'FATOR CONVERSÃO', 'ESTOQUE'])
        sheet.append([456, 'Farinha', 25, 50, 'KG', 888])
        sheet = book.create_sheet('EMBALAGEM')
        sheet.append(['CÓDIGO', 'DESCRIÇÃO PRODUTO', 'BCS', 'SJN', 'MEP', 'FCL 1', 'FCL 2', 'FCL 3', 'FCL 4', 'ESTOQUE MINIMO', 'ESTOQUE MÁXIMO'])
        sheet.append([789, 'Bandeja', 3, 1, 2, None, 0.5, 0.5, 0.5, '=SUM(C2:I2)', '=ROUNDUP(J2*1.2,0)'])
        return book

    def extract(self, book):
        buffer = io.BytesIO()
        book.save(buffer)
        book.close()
        return extract_limits(buffer.getvalue())

    def test_quantity_headers_not_x_markers_bakery_scope_and_fractional_rounding(self):
        data, report = self.extract(self.workbook())
        pairs = {(pair['id'], pair['store']): (pair['estoque_minimo'], pair['estoque_maximo']) for pair in data['limits']}
        self.assertEqual(pairs['123', 'BCS'], (0, 7))
        self.assertEqual(pairs['123', 'SJN'], (0, 2))
        self.assertEqual(pairs['456', 'BCS'], (25, 50))
        self.assertNotIn(('456', 'SJN'), pairs)
        self.assertEqual(pairs['789', 'FCL2'], (0.5, 1))
        self.assertEqual(pairs['789', 'BCS'], (3, 4))
        self.assertFalse(report['pending'])
        self.assertNotIn('999', json.dumps(data))

    def test_manual_multi_store_override_is_quarantined(self):
        book = self.workbook()
        book['EMBALAGEM']['J2'] = 7
        data, report = self.extract(book)
        self.assertFalse(any(pair['id'] == '789' for pair in data['limits']))
        self.assertNotIn('789', data['scope'])
        self.assertEqual(report['pending_products'], ['789'])


if __name__ == '__main__':
    unittest.main()
