"""Review the store columns of the supplied workbook, without evaluating Excel.

The layouts and replenishment formulas below belong to USO E CONSUMO (1).xlsx.
Bakery thresholds are assigned to BCS as explicitly confirmed by the owner.
This command generates reviewable data; it never connects to the database.
"""
import argparse
from collections import Counter
from decimal import Decimal, ROUND_CEILING
import hashlib
import io
import json
from pathlib import Path

import openpyxl

from .catalog_csv import analyze_csv, decimal_cell, normalized
from .import_workbook import ALIASES, code_text, extract_catalog
from .stock_limits import WORKBOOK_STORES as STORES

RULES = {
    'embalagem': 'minimum_plus_20_percent',
    'acougue': 'minimum_times_two',
    'sacolasupermais': 'minimum_times_two',
    'limpeza': 'maximum_only', 'consumogeral': 'maximum_only',
    'escritorio': 'maximum_only', 'bobinapdv': 'maximum_only',
    'etiquetas': 'maximum_only', 'bobinaestrela': 'maximum_only',
    'padaria': 'bakery_bcs', 'consumointernopadaria': 'bakery_bcs',
}


def extract_limits(content):
    source, catalog_report = extract_catalog(content)
    products = {item['id']: item for item in analyze_csv(source)['rows']}
    report = {'source_sha256': hashlib.sha256(content).hexdigest(),
              'conflicting_codes': [item['id'] for item in catalog_report['conflicts']],
              'warnings': [], 'pending': [], 'sheets': []}
    candidates = {}
    scope = set()
    blocked = set()
    values = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    formulas = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in values:
            rule = RULES.get(normalized(sheet.title))
            if not rule or sheet.sheet_state != 'visible':
                continue
            iterator = sheet.iter_rows(max_col=100)
            formula_iterator = formulas[sheet.title].iter_rows(max_col=100)
            mapping = None
            for row_number, (row, formula_row) in enumerate(zip(iterator, formula_iterator), 1):
                names = [normalized(str(cell.value or '')) for cell in row]
                mapping = {key: next((i for i, name in enumerate(names) if name in aliases), None)
                           for key, aliases in ALIASES.items()}
                if mapping['id'] is not None and mapping['nome'] is not None:
                    break
                if row_number >= 15:
                    mapping = None
                    break
            if mapping is None:
                continue
            # LIMPEZA has an earlier group of store names containing X markers.
            # Only quantity headers after the description are considered; the
            # rightmost matching header is the actual numeric store column.
            columns = {store: next((i for i in range(len(names)-1, mapping['nome'], -1)
                                   if names[i] in (normalized(store), normalized(store)+'2')), None)
                       for store in STORES}
            if rule != 'bakery_bcs' and any(index is None for index in columns.values()):
                raise ValueError(f'Colunas de lojas incompletas na aba {sheet.title}.')
            stats = {'sheet': sheet.title, 'rule': rule, 'rows': 0}
            report['sheets'].append(stats)
            for number, (row, formula_row) in enumerate(zip(iterator, formula_iterator), row_number+1):
                if number > 20000:
                    raise ValueError('A aba excede o limite de 20 mil linhas.')
                pid = code_text(row[mapping['id']])
                if pid not in products:
                    continue
                stats['rows'] += 1
                scope.add(pid)
                provenance = {'id': pid, 'sheet': sheet.title, 'row': number}

                def numeric(cell):
                    if cell.value is None or str(cell.value).strip() == '':
                        return None
                    if isinstance(cell.value, bool):
                        raise ValueError('Quantidade booleana inválida.')
                    return decimal_cell(str(cell.value), 'estoque_maximo')

                try:
                    global_values = {key: numeric(row[mapping[key]]) if mapping[key] is not None else None
                                     for key in ('estoque_minimo', 'estoque_maximo')}
                    quotas = {'BCS': None} if rule == 'bakery_bcs' else {
                        store: numeric(row[index]) for store, index in columns.items()}
                    positive = {store: value for store, value in quotas.items() if value is not None and value > 0}
                    if rule == 'bakery_bcs':
                        if all(value is None for value in global_values.values()):
                            continue
                        pairs = [('BCS', global_values['estoque_minimo'] or 0, global_values['estoque_maximo'] or 0)]
                    elif not positive:
                        continue
                    elif rule == 'maximum_only':
                        # Preserve the separate store quantities even if an
                        # aggregate was manually overridden in the spreadsheet.
                        aggregate = global_values['estoque_maximo']
                        total = sum(positive.values())
                        if aggregate is not None and aggregate != total:
                            report['warnings'].append({**provenance, 'message': 'Máximo agregado difere da soma; usadas as colunas individuais das lojas.'})
                        if aggregate == 0:
                            raise ValueError('Máximo agregado zero contradiz as quantidades por loja; revisar item descontinuado.')
                        pairs = [(store, 0, quota) for store, quota in positive.items()]
                    else:
                        multiplier = Decimal('1.2') if rule == 'minimum_plus_20_percent' else Decimal(2)

                        def maximum(minimum):
                            result = Decimal(str(minimum))*multiplier
                            return float(result.to_integral_value(rounding=ROUND_CEILING)) if multiplier == Decimal('1.2') else float(result)

                        pairs = [(store, quota, maximum(quota)) for store, quota in positive.items()]
                        # A literal override can be attributed without guessing
                        # only when the row has exactly one destination store.
                        literals = {key: global_values[key] for key in global_values
                                    if mapping[key] is not None and formula_row[mapping[key]].data_type != 'f'
                                    and global_values[key] is not None}
                        if len(positive) == 1:
                            store, minimum, limit = pairs[0]
                            minimum = literals.get('estoque_minimo', minimum)
                            limit = literals.get('estoque_maximo', maximum(minimum))
                            pairs = [(store, minimum, limit)]
                        elif literals:
                            raise ValueError('Limite agregado manual com várias lojas; não foi distribuído automaticamente.')
                    for store, minimum, maximum_value in pairs:
                        if maximum_value and minimum > maximum_value:
                            raise ValueError('Mínimo maior que máximo.')
                        pair = {'id': pid, 'store': store, 'estoque_minimo': minimum,
                                'estoque_maximo': maximum_value, 'sheet': sheet.title, 'row': number,
                                'rule': rule, 'unidade': products[pid].get('unidade', 'UN')}
                        key = (pid, store)
                        previous = candidates.get(key)
                        if previous and any(previous[field] != pair[field] for field in ('estoque_minimo', 'estoque_maximo', 'unidade')):
                            raise ValueError('Limites conflitantes para o mesmo produto e loja.')
                        candidates[key] = pair
                except ValueError as error:
                    blocked.add(pid)
                    report['pending'].append({**provenance, 'message': str(error)})
    finally:
        values.close()
        formulas.close()
    limits = sorted((pair for pair in candidates.values() if pair['id'] not in blocked), key=lambda pair: (pair['id'], pair['store']))
    report['limits'] = len(limits)
    report['products_with_limits'] = len({pair['id'] for pair in limits})
    report['by_store'] = dict(sorted(Counter(pair['store'] for pair in limits).items()))
    report['pending_products'] = sorted(blocked)
    data = {'version': 'limites_excel_lojas_20261019_v1', 'source_sha256': report['source_sha256'],
            'policy': {'maximum_only_minimum': 0, 'bakery_store': 'BCS',
                       'embalagem': 'Mínimo da loja; máximo arredondado para cima em 120%, respeitando ajustes literais de loja única.',
                       'acougue': 'Mínimo da loja; máximo igual ao dobro do mínimo.',
                       'empty_cells': 'Sem limite atribuído; saldos físicos não são importados.'},
            'scope': sorted(scope-blocked), 'limits': limits}
    return data, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--report', type=Path, required=True)
    args = parser.parse_args()
    data, report = extract_limits(args.source.read_bytes())
    for path, value in ((args.output, data), (args.report, report)):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(f'{report["limits"]} limites por loja; {len(report["pending_products"])} produtos para revisão.')


if __name__ == '__main__':
    main()
